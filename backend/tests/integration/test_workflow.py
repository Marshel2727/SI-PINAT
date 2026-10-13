import io
import uuid
from sqlalchemy import inspect
from app.repositories.store import one, rows
from openpyxl import load_workbook

def create(client,mode="DEMO",preference="ES-01"):
    body={"mode":mode,"semester":4,"preferred_specialization_code":preference}
    if mode=="FULL":
        body.update(nim=uuid.uuid4().hex[:20],full_name="Test Participant",class_name="TEST")
    response=client.post("/api/v1/assessments",json=body)
    assert response.status_code==201,response.text
    return "/api/v1/assessments/"+response.json()["data"]["public_id"]

def grades(client,base):
    master=client.get(base+"/courses").json()["data"]
    package=[{"semester":c["semester"],"kode_mk":c["code"],"nama_mk":c["name"],"nilai_huruf":"A"} for c in master["courses"]]
    response=client.put(base+"/grades",json={"grades":package})
    assert response.status_code==200,response.text
    assert response.json()["data"]["ready"]
    return package

def correct_options(db,packet):
    return {item["id"]:next(o["id"] for o in rows(db,"question_options",question_id=int(item["question_id"])) if o["is_correct"]) for item in packet}

def test_schema_and_composite_foreign_keys(prepared_database):
    inspector=inspect(prepared_database)
    assert len([t for t in inspector.get_table_names() if t!="alembic_version"])==23
    assert sum(len(inspector.get_foreign_keys(t)) for t in inspector.get_table_names())==34
    foreign=inspector.get_foreign_keys("assessment_answers")
    assert {tuple(f["constrained_columns"]) for f in foreign}=={("assessment_item_id","question_id"),("question_id","option_id")}

def test_demo_complete_and_idempotence(client,prepared_database):
    base=create(client)
    grades(client,base)
    response=client.post(base+"/test/start")
    assert response.status_code==200,response.text
    packet=response.json()["data"]
    assert len(packet)==30
    assert all("is_correct" not in o for i in packet for o in i["options"])
    assert client.post(base+"/complete").status_code==409
    with prepared_database.connect() as db:
        options=correct_options(db,packet)
    for item in packet:
        if item["is_representative"]:
            assert client.put(base+"/test/items/"+item["id"]+"/answer",json={"option_id":options[item["id"]]}).status_code==200
    assert client.post(base+"/demo/fill-examples").status_code==200
    result=client.post(base+"/complete")
    assert result.status_code==200,result.text
    value=result.json()["data"]
    assert len(value["scores"])==3
    assert value["answer_sources"]=={"USER":3,"EXAMPLE":27}
    again=client.post(base+"/complete").json()["data"]
    assert again["run_id"]==value["run_id"]
    trace=client.get(base+"/result/trace").json()["data"]
    assert trace["centroid_denominator"]>0
    assert client.put(base+"/grades",json={"grades":[]}).status_code==409
    pdf=client.get(base+"/result/report.pdf")
    assert pdf.status_code==200 and pdf.content.startswith(b"%PDF")
    # UUID alone cannot disclose result or report.
    from fastapi.testclient import TestClient
    from app.main import app
    outsider=TestClient(app)
    assert outsider.get(base+"/result").status_code==404
    assert outsider.get(base+"/result/report.pdf").status_code==404

def test_full_rejects_example_wrong_option_and_incomplete(client,prepared_database):
    base=create(client,"FULL")
    grades(client,base)
    packet=client.post(base+"/test/start").json()["data"]
    assert client.post(base+"/demo/fill-examples").status_code==403
    wrong=packet[1]["options"][0]["id"]
    assert client.put(base+"/test/items/"+packet[0]["id"]+"/answer",json={"option_id":wrong}).status_code==422
    assert client.put(base+"/test/items/"+packet[0]["id"]+"/answer",json={"option_id":packet[0]["options"][0]["id"]}).status_code==200
    assert client.post(base+"/complete").status_code==422

def test_invalid_import_preserves_grades_and_csrf_admin(client):
    base=create(client)
    package=grades(client,base)
    response=client.post(base+"/grades/import",files={"file":("bad.csv",b"bad,header\n1,2","text/csv")},data={"confirm":"true"})
    assert response.status_code==422
    assert len(client.get(base+"/grades").json()["data"]["grades"])==len(package)
    assert client.get("/api/v1/admin/dashboard").status_code==401
    assert client.post("/api/v1/auth/login",json={"email":"test@example.invalid","password":"test-only-password"},headers={"X-CSRF-Token":"invalid"}).status_code==403
    assert client.post("/api/v1/auth/login",json={"email":"test@example.invalid","password":"test-only-password"}).status_code==200
    export=client.get("/api/v1/admin/reports/assessments.xlsx?mode=DEMO")
    assert export.status_code==200
    workbook=load_workbook(io.BytesIO(export.content),read_only=True)
    assert workbook.active["A1"].value=="Sesi"
    assert all(row[1]=="DEMO" for row in workbook.active.iter_rows(min_row=2,values_only=True))
    assert client.post("/api/v1/auth/logout").status_code==204
    assert client.get("/api/v1/auth/me").status_code==401

def test_fuzzy_preference_invariance(client,prepared_database):
    results=[]
    for preference in ["ES-01","SC-02"]:
        base=create(client,"FULL",preference)
        grades(client,base)
        packet=client.post(base+"/test/start").json()["data"]
        with prepared_database.connect() as db:
            options=correct_options(db,packet)
        for item in packet:
            assert client.put(base+"/test/items/"+item["id"]+"/answer",json={"option_id":options[item["id"]]}).status_code==200
        value=client.post(base+"/complete")
        assert value.status_code==200,value.text
        result=value.json()["data"]
        results.append([(s["specialization_code"],s["suitability_score"],s["rank_position"]) for s in result["scores"]])
        assert result["has_tie"] and len(result["top_specialization_codes"])==3
    assert results[0]==results[1]
