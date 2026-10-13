import csv
import hashlib
import io
import uuid
import zipfile
from pathlib import Path
from openpyxl import load_workbook, Workbook
from app.core.errors import fail
from app.core.config import settings
from app.repositories.store import add, safe
from app.services.assessments import editable
from app.services.courses import for_assessment
from app.services.assessment_grades import check, readiness, save

HEADERS=["semester","kode_mk","nama_mk","nilai_huruf"]

def parse(filename, content, content_type):
    if not content or len(content)>10*1024*1024:
        raise ValueError("Berkas kosong atau lebih dari 10 MB.")
    suffix=Path(filename).suffix.lower()
    if suffix not in {".csv",".xlsx"}:
        raise ValueError("Hanya berkas CSV dan XLSX yang didukung.")
    if suffix==".csv":
        if content_type not in {"text/csv","text/plain","application/vnd.ms-excel","application/octet-stream"}:
            raise ValueError("Tipe berkas CSV tidak sesuai.")
        try:
            values=list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        except (UnicodeError,csv.Error):
            raise ValueError("CSV harus menggunakan UTF-8.")
    else:
        if content_type not in {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","application/octet-stream"} or not zipfile.is_zipfile(io.BytesIO(content)):
            raise ValueError("Struktur XLSX tidak sah.")
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries=archive.infolist()
            if len(entries)>1000 or sum(v.file_size for v in entries)>40*1024*1024:
                raise ValueError("Ukuran dekompresi XLSX terlalu besar.")
            if any("vbaproject" in e.filename.lower() or ".." in Path(e.filename).parts or e.filename.startswith(("/","\\")) for e in entries):
                raise ValueError("Makro atau struktur arsip tidak diizinkan.")
            if "[Content_Types].xml" not in archive.namelist() or "xl/workbook.xml" not in archive.namelist():
                raise ValueError("Berkas bukan workbook XLSX.")
        try:
            workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
            sheet=workbook.active
            values=[]
            for index,row in enumerate(sheet.iter_rows()):
                if index>1000:
                    raise ValueError("Maksimal 1.000 baris data.")
                if any(cell.data_type=="f" for cell in row):
                    raise ValueError("Formula tidak diizinkan.")
                values.append([cell.value if cell.value is not None else "" for cell in row])
            workbook.close()
        except ValueError:
            raise
        except Exception:
            raise ValueError("Workbook XLSX tidak dapat dibaca.")
    if not values or [str(c).strip() for c in values[0]]!=HEADERS:
        raise ValueError("Header wajib: "+", ".join(HEADERS)+".")
    data=[]
    for row in values[1:]:
        if not any(str(c).strip() for c in row):
            continue
        if len(row)!=4:
            raise ValueError("Setiap baris wajib mempunyai empat kolom.")
        if any(str(c).lstrip().startswith(("=","+","-","@")) for c in row):
            raise ValueError("Formula atau awalan berbahaya tidak diizinkan.")
        data.append(dict(zip(HEADERS,[str(c).strip() for c in row])))
        if len(data)>1000:
            raise ValueError("Maksimal 1.000 baris data.")
    if not data:
        raise ValueError("Berkas belum memiliki nilai.")
    return data

def import_file(db, assessment, filename, content, content_type, confirm):
    editable(assessment)
    filename=Path(filename or "upload").name[:255]
    error=[]
    try:
        data=parse(filename,content,content_type)
        grades,error=check(db,assessment,data)
        state=readiness(db,assessment,grades)
        if not state["ready"]:
            error.append({"missing_required":state["missing_required"],"messages":state["errors"]})
    except ValueError as exc:
        data,grades,state=[],[],{"ready":False}
        error=[{"row":None,"messages":[str(exc)]}]
    key=uuid.uuid4().hex+Path(filename).suffix.lower() if not error else None
    if key:
        directory=Path(settings().upload_dir)
        directory.mkdir(parents=True,exist_ok=True)
        (directory/key).write_bytes(content)
    identifier=add(db,"transcript_imports",assessment_id=assessment["id"],original_filename=filename,
                   storage_key=key,file_hash=hashlib.sha256(content).hexdigest(),
                   status="INVALID" if error else "VALID",validation_errors=error or None)
    if error:
        # Retain invalid import history, preserve previous valid grades.
        db.commit()
        fail(422,"Impor ditolak. Nilai sebelumnya tetap tersimpan.","IMPORT_INVALID",error)
    result=save(db,assessment,data,"IMPORT",False) if confirm else {"grades":grades,**state}
    return {"import_id":identifier,"committed":confirm,"preview":safe(data),**result}

def template(db, assessment, format):
    records=[dict(zip(HEADERS,[c["semester"],c["code"],c["name"],""])) for c in for_assessment(db,assessment)["courses"]]
    if format=="csv":
        output=io.StringIO()
        writer=csv.DictWriter(output,fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows([{k:export_text(v) for k,v in row.items()} for row in records])
        return ("\ufeff"+output.getvalue()).encode("utf-8"),"text/csv; charset=utf-8","template-nilai.csv"
    if format!="xlsx":
        fail(422,"Format template harus csv atau xlsx.")
    workbook=Workbook()
    sheet=workbook.active
    sheet.title="Nilai"
    sheet.append(HEADERS)
    for row in records:
        sheet.append([export_text(row[key]) for key in HEADERS])
    output=io.BytesIO()
    workbook.save(output)
    return output.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","template-nilai.xlsx"

def export_text(value):
    if isinstance(value,str) and value.lstrip().startswith(("=","+","-","@")):
        return "'"+value
    return value
